import io
from datetime import datetime

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.security import get_current_user
from app.models.models import Task, User

router = APIRouter(prefix="/api/export", tags=["export"])

STATUS_LABELS = {
    "NOVA": "Nova",
    "PENDENTE": "Pendente",
    "EM_ANALISE": "Em análise",
    "EM_EXECUCAO": "Em andamento",
    "AGUARDANDO_TERCEIRO": "Aguardando terceiro",
    "BLOQUEADA": "Bloqueada",
    "EM_VALIDACAO": "Em validação",
    "CONCLUIDA": "Concluída",
    "CANCELADA": "Cancelada",
}

PRIORITY_LABELS = {
    "BAIXA": "Baixa",
    "MEDIA": "Média",
    "ALTA": "Alta",
    "CRITICA": "Crítica",
}

HEADERS = [
    "ID", "Atividade", "Tipo", "Prioridade", "Status",
    "Solicitante", "Marca", "Responsável", "Departamento",
    "Dt. Solicitação", "Dt. Entrega", "Motivo Bloqueio", "Observações",
]


def _get_tasks(db: Session, current_user: User, status: str | None, responsible_id: int | None, marca: str | None):
    query = db.query(Task)
    if status:
        query = query.filter(Task.status == status)
    if marca:
        query = query.filter(Task.marca == marca)
    if current_user.role != "GESTOR":
        query = query.filter(Task.responsible_id == current_user.id)
    elif responsible_id:
        query = query.filter(Task.responsible_id == responsible_id)
    return query.order_by(Task.updated_at.desc()).all()


def _task_row(task: Task) -> list:
    return [
        task.id,
        task.title or "",
        task.task_type or "",
        PRIORITY_LABELS.get(task.priority, task.priority or ""),
        STATUS_LABELS.get(task.status, task.status or ""),
        task.requester or "",
        task.marca or "",
        task.responsible_user.name if task.responsible_user else "",
        task.department.name if task.department else "",
        task.request_date or "",
        task.due_date or "",
        task.blocking_reason or "",
        task.observations or "",
    ]


@router.get("/tasks/excel")
def export_excel(
    status: str | None = Query(default=None),
    responsible_id: int | None = Query(default=None),
    marca: str | None = Query(default=None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    tasks = _get_tasks(db, current_user, status, responsible_id, marca)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Demandas"

    header_fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=10)
    alt_fill = PatternFill(start_color="EEF2F7", end_color="EEF2F7", fill_type="solid")

    for col, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx, task in enumerate(tasks, 2):
        for col, value in enumerate(_task_row(task), 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if row_idx % 2 == 0:
            for col in range(1, len(HEADERS) + 1):
                ws.cell(row=row_idx, column=col).fill = alt_fill

    col_widths = [6, 40, 14, 12, 22, 20, 14, 20, 22, 14, 14, 30, 30]
    for col, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"demandas_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/tasks/pdf")
def export_pdf(
    status: str | None = Query(default=None),
    responsible_id: int | None = Query(default=None),
    marca: str | None = Query(default=None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    tasks = _get_tasks(db, current_user, status, responsible_id, marca)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=1 * cm,
        rightMargin=1 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "Title",
        parent=styles["Heading1"],
        fontSize=13,
        textColor=colors.HexColor("#1E3A5F"),
        spaceAfter=8,
    )
    cell_style = ParagraphStyle("Cell", parent=styles["Normal"], fontSize=7, leading=9)
    header_style = ParagraphStyle(
        "Header",
        parent=styles["Normal"],
        fontSize=7.5,
        textColor=colors.white,
        fontName="Helvetica-Bold",
        leading=10,
    )

    header_row = [Paragraph(h, header_style) for h in HEADERS]
    rows = [header_row] + [
        [Paragraph(str(v), cell_style) for v in _task_row(t)] for t in tasks
    ]

    col_widths = [
        0.9 * cm, 5 * cm, 2 * cm, 1.8 * cm, 2.5 * cm,
        2.4 * cm, 1.8 * cm, 2.4 * cm, 2.4 * cm,
        1.8 * cm, 1.8 * cm, 3.2 * cm, 3.2 * cm,
    ]

    table = Table(rows, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A5F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EEF2F7")]),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ]))

    title_text = f"Demandas TaskFlow — {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    if status:
        title_text += f"  |  Status: {STATUS_LABELS.get(status, status)}"
    if responsible_id and tasks:
        responsible_name = tasks[0].responsible_user.name if tasks[0].responsible_user else str(responsible_id)
        title_text += f"  |  Responsável: {responsible_name}"
    if marca:
        title_text += f"  |  Marca: {marca}"

    doc.build([Paragraph(title_text, title_style), Spacer(1, 0.3 * cm), table])
    buf.seek(0)

    filename = f"demandas_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
